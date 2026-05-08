import fitz  # PyMuPDF
from pptx import Presentation
import os


# -----------------------------
# PDF TEXT EXTRACTION (FAST)
# -----------------------------
def extract_pdf_text(file_path: str) -> str:
    doc = fitz.open(file_path)
    text = ""

    for page in doc:
        text += page.get_text()

    doc.close()
    return text


# -----------------------------
# PPT TEXT EXTRACTION
# -----------------------------
def extract_ppt_text(file_path: str) -> str:
    prs = Presentation(file_path)
    text = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)

    return " ".join(text)


# -----------------------------
# PDF IMAGE EXTRACTION
# -----------------------------
def extract_images_from_pdf(file_path: str, output_dir: str):
    doc = fitz.open(file_path)
    os.makedirs(output_dir, exist_ok=True)

    image_paths = []

    for page_index in range(len(doc)):
        page = doc[page_index]
        images = page.get_images(full=True)

        for img_index, img in enumerate(images):
            xref = img[0]
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]

            img_path = os.path.join(
                output_dir,
                f"pdf_img_{page_index}_{img_index}.png"
            )

            with open(img_path, "wb") as f:
                f.write(image_bytes)

            image_paths.append(img_path)

    doc.close()
    return image_paths


# -----------------------------
# PPT IMAGE EXTRACTION
# -----------------------------
def extract_images_from_ppt(file_path: str, output_dir: str):
    prs = Presentation(file_path)
    os.makedirs(output_dir, exist_ok=True)

    image_paths = []

    for slide_index, slide in enumerate(prs.slides):
        for shape_index, shape in enumerate(slide.shapes):

            try:
                image = shape.image
                image_bytes = image.blob

                img_path = os.path.join(
                    output_dir,
                    f"ppt_img_{slide_index}_{shape_index}.png"
                )

                with open(img_path, "wb") as f:
                    f.write(image_bytes)

                image_paths.append(img_path)

            except Exception:
                continue

    return image_paths


# -----------------------------
# PPT TABLE EXTRACTION (FAST)
# -----------------------------
def extract_tables_from_ppt(file_path: str) -> str:
    prs = Presentation(file_path)

    table_texts = []

    for slide_index, slide in enumerate(prs.slides):
        for shape in slide.shapes:

            if shape.has_table:
                table = shape.table

                table_str = f"\n--- PPT TABLE (Slide {slide_index+1}) ---\n"

                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    table_str += " | ".join(row_text) + "\n"

                table_texts.append(table_str)

    return "\n\n".join(table_texts)